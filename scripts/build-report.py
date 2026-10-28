"""Generate the approved engineering DOCX and PDF from shared source/evidence.

This editorial tool combines report-source.json, the frozen local validation
checkpoint and approved screenshots into two parallel nine-page layouts. It also
regenerates the architecture asset, removes private/exact-date metadata, and writes
files under docs plus scratch files under .local/report-layout. It is deliberately
not part of ordinary Vite builds: those copy the already reviewed report bytes.

Requires Pillow, pypdf, python-docx, ReportLab and pdftoppm (or COBALT_PDFTOPPM).
Regeneration requires full visual review of both formats before publication. The
report describes a browser HMI simulation, not live acquisition, certified controls
or measured hardware response; evidence tokens must retain their stated limits.
"""
import json
import os
import re
import subprocess
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image as PdfImage
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Polygon
from reportlab.graphics import renderPDF, renderSVG

# Resolve all editorial inputs/outputs relative to the repository, independent of shell cwd.
ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
source = json.loads(Path('docs/report-source.json').read_text(encoding='utf-8'))
evidence = json.loads(Path('docs/validation-evidence.json').read_text(encoding='utf-8'))
scratch = Path('.local/report-layout')
scratch.mkdir(parents=True, exist_ok=True)
scenario = {s['id']: s for s in evidence['scenarios']}
normal = scenario['normal-ramp']['finalTelemetry']
js = [b for b in evidence['bundles'] if b['file'].endswith('.js')]
css = [b for b in evidence['bundles'] if b['file'].endswith('.css')]
# Fill measurements from the selected checkpoint; never invent fresh counts or timings.
tokens = {
    'unitPassed': evidence['unit']['passed'], 'browserPassed': evidence['browserTests']['passed'],
    'normalRPM': f"{normal['rpm']:.1f}", 'normalLoad': f"{normal['engineLoad']:.1f}",
    'normalOil': f"{normal['oilPressure']:.1f}", 'normalCoolant': f"{normal['coolantTemp']:.1f}",
    'overheatCoolant': f"{scenario['overheat']['finalTelemetry']['coolantTemp']:.1f}",
    'jsBytes': f"{sum(b['bytes'] for b in js):,}", 'jsGzipBytes': f"{sum(b['gzipBytes'] for b in js):,}",
    'cssBytes': f"{sum(b['bytes'] for b in css):,}", 'wallMs': evidence['browser']['wallMillisecondsToObserve'],
    'browserDuration': f"{evidence['browserTests']['durationMs'] / 1000:.1f}",
    'unitDuration': f"{evidence['unit']['durationMs'] / 1000:.2f}",
}
def fill(text):
    """Replace report tokens with recorded measurements; unknown tokens surface as errors."""
    return re.sub(r'\{\{(\w+)\}\}', lambda m: str(tokens[m.group(1)]), text)

def architecture():
    """Draw the simulated command-to-display data flow as SVG and report-ready raster."""
    drawing = Drawing(660, 190)
    def node(x, y, width, label, detail):
        drawing.add(Rect(x, y, width, 43, rx=5, ry=5, fillColor=colors.HexColor('#F0F5FC'), strokeColor=colors.HexColor('#9DB5D2'), strokeWidth=0.8))
        drawing.add(String(x + width / 2, y + 25, label, textAnchor='middle', fontName='Helvetica-Bold', fontSize=11, fillColor=colors.HexColor('#112442')))
        drawing.add(String(x + width / 2, y + 11, detail, textAnchor='middle', fontName='Helvetica', fontSize=9, fillColor=colors.HexColor('#475D78')))
    def arrow(x1, y1, x2, y2):
        drawing.add(Line(x1, y1, x2, y2, strokeColor=colors.HexColor('#175CD3'), strokeWidth=1.4))
        if x2 > x1: points = [x2, y2, x2 - 5, y2 + 3, x2 - 5, y2 - 3]
        else: points = [x2, y2, x2 - 3, y2 + 5, x2 + 3, y2 + 5]
        drawing.add(Polygon(points, fillColor=colors.HexColor('#175CD3'), strokeColor=None))
    for args in [(0, 110, 116, 'Commands', 'mode and scenario'), (136, 110, 116, 'Engine model', 'seed and fixed step'), (272, 110, 116, 'Validation', 'version, bounds, time'), (408, 110, 116, 'Received state', 'last accepted sample'), (544, 110, 116, 'Instruments', 'gauges and preview')]: node(*args)
    for x in [116, 252, 388, 524]: arrow(x, 131, x + 20, 131)
    node(136, 22, 116, 'Settings', 'validated rules')
    node(272, 22, 116, 'Diagnostics', 'alerts and events')
    node(408, 22, 116, 'History', 'bounded trend')
    arrow(252, 43, 272, 43)
    arrow(466, 110, 466, 65)
    drawing.add(Line(440, 110, 440, 81, strokeColor=colors.HexColor('#175CD3'), strokeWidth=1.4))
    drawing.add(Line(440, 81, 330, 81, strokeColor=colors.HexColor('#175CD3'), strokeWidth=1.4))
    arrow(330, 81, 330, 65)
    renderSVG.drawToFile(drawing, 'docs/assets/architecture.svg')
    renderPDF.drawToFile(drawing, str(scratch / 'architecture.pdf'))
    executable = os.environ.get('COBALT_PDFTOPPM', 'pdftoppm')
    subprocess.run([executable, '-png', '-singlefile', '-scale-to', '1800', str(scratch / 'architecture.pdf'), str(scratch / 'architecture')], check=True, capture_output=True)
    return scratch / 'architecture.png'

# Generate the shared illustration once so Word/PDF explain the same data relationships.
architecture_path = architecture()
def image_file(block):
    """Resolve approved images and optional report crops without changing the original capture."""
    if block['kind'] == 'architecture': return architecture_path
    path = Path(block['path'])
    if block.get('cropHeight'):
        with Image.open(path) as original:
            cropped = original.crop((0, 0, original.width, min(block['cropHeight'], original.height)))
            path = scratch / 'dashboard-overview.png'
            cropped.save(path)
    return path
def dimensions(block, path):
    """Fit images proportionally inside the report page bounds to avoid clipping."""
    with Image.open(path) as img: width, height = img.size
    max_w = 6.7
    max_h = block.get('height', 2.05)
    scale = min(max_w / width, max_h / height)
    return width * scale, height * scale

# Word layout: page geometry, readable typography and cleared inherited decorative borders.
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.8)
section.footer_distance = Inches(0.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = True
doc.styles['Caption'].font.bold = False
doc.styles['Caption'].paragraph_format.space_after = Pt(10)
for style in doc.styles:
    paragraph_properties = style._element.find(qn('w:pPr'))
    if paragraph_properties is not None:
        for border in list(paragraph_properties.findall(qn('w:pBdr'))):
            paragraph_properties.remove(border)
# Keep approved visible attribution in the body while stripping incidental personal metadata.
doc.core_properties.author = ''
# Omit precise metadata dates; the report's visible edition is Fall 2026.
for date_tag in ['dcterms:created', 'dcterms:modified', 'cp:lastPrinted']:
    for element in list(doc.core_properties._element.findall(qn(date_tag))):
        doc.core_properties._element.remove(element)
doc.core_properties.revision = 1
doc.core_properties.last_modified_by = ''
doc.core_properties.title = f"{source['title']} Engineering Report"
doc.core_properties.subject = 'Simulated instrumentation architecture and validation'
doc.core_properties.comments = ''
doc.core_properties.keywords = 'simulation, instrumentation, diagnostics'
# Stable version/page footer supports navigation during a printed or exported report demo.
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run(f"Cobalt Instrument Studio {source['version']}  |  ").font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

# PDF styles mirror the Word hierarchy while remaining independent layout definitions.
pdf_body = ParagraphStyle('body', fontName='Helvetica', fontSize=10.7, leading=14.1, textColor=colors.black, spaceAfter=9)
pdf_title = ParagraphStyle('title', parent=pdf_body, fontName='Helvetica-Bold', fontSize=25, leading=30, spaceAfter=10)
pdf_subtitle = ParagraphStyle('subtitle', parent=pdf_body, fontSize=14, leading=19, spaceAfter=10)
pdf_heading = ParagraphStyle('h1', parent=pdf_body, fontName='Helvetica-Bold', fontSize=16, leading=20, spaceAfter=12)
pdf_h2 = ParagraphStyle('h2', parent=pdf_body, fontName='Helvetica-Bold', fontSize=12, leading=16, spaceBefore=7, spaceAfter=6)
pdf_caption = ParagraphStyle('caption', parent=pdf_body, fontName='Helvetica-Oblique', fontSize=8.8, leading=11.2, spaceAfter=12)
pdf_cell = ParagraphStyle('cell', parent=pdf_body, fontSize=9.2, leading=12, spaceAfter=0)
pdf_headcell = ParagraphStyle('headcell', parent=pdf_cell, fontName='Helvetica-Bold', textColor=colors.white)
pdf_story = []
# Escape authored text before handing it to ReportLab markup parsing.
def para(text): return Paragraph(escape(fill(text)), pdf_body)

def add_repository_link(paragraph, url):
    """Attach the approved public repository hyperlink using Word relationship metadata."""
    paragraph.add_run(' | ').font.size = Pt(10)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '175CD3'); properties.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '20'); properties.append(size)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); properties.append(underline)
    run.append(properties)
    text = OxmlElement('w:t'); text.text = 'GitHub'; run.append(text)
    link.append(run); paragraph._p.append(link)

def add_table(block):
    """Emit matching semantic tables in both formats, repeating headers across pages."""
    rows = [block['columns']] + block['rows']
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, block['widths']): column.width = Inches(width)
    for row_index, values in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for cell_index, value in enumerate(values):
            cell = row.cells[cell_index]
            cell.width = Inches(block['widths'][cell_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '173B68' if row_index == 0 else 'F1F5FB' if row_index % 2 else 'FFFFFF')
            props.append(shading)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                item = OxmlElement(f'w:{side}'); item.set(qn('w:w'), '80'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement(f'w:{side}'); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            props.append(borders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(fill(str(value)))
            run.font.size = Pt(9.5)
            run.font.name = 'Arial'
            run.bold = row_index == 0
            run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)
            if cell_index == len(values) - 1 and len(values) == 3 and block['columns'][-1] == 'Status': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    pdf_rows = [[Paragraph(escape(fill(str(value))), pdf_headcell if i == 0 else pdf_cell) for value in row] for i, row in enumerate(rows)]
    pdf_table = Table(pdf_rows, colWidths=[w * 72 for w in block['widths']], repeatRows=1, hAlign='CENTER')
    commands = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#173B68')), ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6), ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6)]
    for i in range(1, len(rows), 2): commands.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F1F5FB')))
    pdf_table.setStyle(TableStyle(commands))
    pdf_story.extend([pdf_table, Spacer(1, 10)])

# Render one shared ordered content model into Word paragraphs and the PDF story.
for index, page in enumerate(source['pages']):
    if index:
        pdf_story.append(PageBreak())
    if index == 0:
        doc.add_paragraph(source['title'], style='Title')
        doc.add_paragraph(source['subtitle'], style='Subtitle')
        credit = doc.add_paragraph(source['authorCredit'])
        credit.paragraph_format.space_after = Pt(6)
        credit.runs[0].font.size = Pt(10)
        credit.runs[0].font.color.rgb = RGBColor.from_string('50617A')
        repository = source.get('authorRepository')
        if repository: add_repository_link(credit, repository)
        doc.add_paragraph(f"Version {source['version']}  |  {source['date']}")
        pdf_credit = ParagraphStyle('credit', parent=pdf_body, fontSize=10, leading=13, textColor=colors.HexColor('#50617A'), spaceAfter=6)
        pdf_credit_text = escape(source['authorCredit']) + (f' | <link href="{escape(repository)}" color="#175CD3"><u>GitHub</u></link>' if repository else '')
        pdf_story.extend([Paragraph(escape(source['title']), pdf_title), Paragraph(escape(source['subtitle']), pdf_subtitle), Paragraph(pdf_credit_text, pdf_credit), para(f"Version {source['version']}  |  {source['date']}")])
    page_heading = doc.add_paragraph(page['title'], style='Heading 1')
    page_heading.paragraph_format.page_break_before = bool(index)
    pdf_story.append(Paragraph(escape(page['title']), pdf_heading))
    for block in page['blocks']:
        kind = block['kind']
        if kind == 'paragraph':
            doc.add_paragraph(fill(block['text'])); pdf_story.append(para(block['text']))
        elif kind == 'heading':
            doc.add_paragraph(block['text'], style='Heading 2'); pdf_story.append(Paragraph(escape(block['text']), pdf_h2))
        elif kind == 'table': add_table(block)
        elif kind in ['image', 'architecture']:
            path = image_file(block); width, height = dimensions(block, path)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
            shape = doc.inline_shapes[-1]
            shape._inline.docPr.set('descr', block['caption'])
            doc.add_paragraph(block['caption'], style='Caption')
            pdf_story.extend([PdfImage(str(path), width=width * 72, height=height * 72, hAlign='CENTER'), Spacer(1, 6), Paragraph(escape(block['caption']), pdf_caption)])

# These stable filenames are consumed by build synchronization, downloads and release assets.
docx_path = Path('docs/Cobalt_Engineering_Report.docx')
pdf_path = Path('docs/Cobalt_Engineering_Report.pdf')
doc.save(docx_path)
def footer_pdf(canvas, document):
    """Add version/page text without leaking formatting state into the page body."""
    canvas.saveState(); canvas.setFont('Helvetica', 8); canvas.setFillColor(colors.HexColor('#50617A'))
    canvas.drawRightString(7.7 * 72, 0.33 * 72, f"Cobalt Instrument Studio {source['version']}  |  {document.page}")
    canvas.restoreState()
pdf = SimpleDocTemplate(str(pdf_path), pagesize=(8.5 * 72, 11 * 72), leftMargin=0.8 * 72, rightMargin=0.8 * 72, topMargin=0.65 * 72, bottomMargin=0.65 * 72, title=f"{source['title']} Engineering Report", author='', subject='Simulated instrumentation architecture and validation', creator='Cobalt report builder')
pdf.build(pdf_story, onFirstPage=footer_pdf, onLaterPages=footer_pdf)
# ReportLab inserts exact creation dates. Remove them from the final artifact.
reader = PdfReader(pdf_path)
writer = PdfWriter(clone_from=reader)
writer.metadata = {
    key: str(value) for key, value in (reader.metadata or {}).items()
    if key not in ['/CreationDate', '/ModDate']
}
with pdf_path.open('wb') as output:
    writer.write(output)
print(f'Created reports: DOCX {docx_path.stat().st_size} bytes; PDF {pdf_path.stat().st_size} bytes')
